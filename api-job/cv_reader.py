from groq import Groq
from dotenv import load_dotenv
from pathlib import Path
import requests
import os
import json
import docx
from PyPDF2 import PdfReader
import base64

load_dotenv()

def read_file(text_file_path):
    context_path = Path(__file__).parent / "template_prompt" / text_file_path
    with open (context_path, "r", encoding="utf-8") as file:
        return file.read()
    
def extract_text_from_pdf(pdf_path):
    reader = PdfReader(pdf_path)
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""
    return text

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    return "\n".join([para.text for para in doc.paragraphs])

# Function to encode the image
def encode_image(image_path):
  with open(image_path, "rb") as image_file:
    return base64.b64encode(image_file.read()).decode('utf-8')

def ask_vision(image_input):
    """
    image_input: chemin (str/Path) ou objet fichier avec .read() (Streamlit UploadedFile, BytesIO) ou bytes.
    Retourne le JSON parsé fourni par le modèle.
    """
    # Lire les bytes de l'image quel que soit le type d'input
    if isinstance(image_input, (str, Path)):
        with open(image_input, "rb") as f:
            image_bytes = f.read()
    elif isinstance(image_input, (bytes, bytearray)):
        image_bytes = bytes(image_input)
    elif hasattr(image_input, "read"):
        # file-like (Streamlit UploadedFile, io.BytesIO, etc.)
        image_bytes = image_input.read()
        # tenter de remettre le curseur au début pour réutilisation éventuelle
        try:
            image_input.seek(0)
        except Exception:
            pass
    else:
        raise ValueError("image_input doit être un chemin, des bytes ou un objet fichier avec .read()")

    base64_image = base64.b64encode(image_bytes).decode("utf-8")

    client = Groq(api_key=os.environ.get("GROQ_API_KEY"))

    chat_completion = client.chat.completions.create(
        messages=[
            {
                "role": "system",
                "content": read_file("contexte_image.txt")
            },
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": read_file("prompt_image.txt")},
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:image/jpeg;base64,{base64_image}",
                        },
                    },
                ],
            }
        ],
        model="meta-llama/llama-4-scout-17b-16e-instruct",
        response_format={"type": "json_object"}
    )

    result = json.loads(chat_completion.choices[0].message.content)
    return result

def read_cv(file_input):
    client = Groq(api_key=os.environ["GROQ_API_KEY"])
    # Si file_input est un chemin (str ou Path)
    if isinstance(file_input, (str, Path)):
        ext = str(file_input).lower()
        if ext.endswith(".pdf"):
            text = extract_text_from_pdf(file_input)
        elif ext.endswith(".docx"):
            text = extract_text_from_docx(file_input)
        else:
            raise ValueError("Le fichier doit être au format PDF ou DOCX.")
    else:
        # Cas UploadedFile (Streamlit)
        name = file_input.name.lower()
        if name.endswith(".pdf"):
            reader = PdfReader(file_input)
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
        elif name.endswith(".docx"):
            doc = docx.Document(file_input)
            text = "\n".join([para.text for para in doc.paragraphs])
        else:
            raise ValueError("Le fichier doit être au format PDF ou DOCX.")

    response = client.chat.completions.create(
        model="openai/gpt-oss-20b",
        messages=[
            {
                "role": "system",
                "content": read_file("context.txt"),
                
            },
            {
                "role": "user",
                "content": f"Analyse le texte ci-dessous (ta réponse doit être dans le format JSON) : {text}",
            }
        ],
        temperature=0.0,
        response_format={"type": "json_object",}
    )
    return response.choices[0].message.content

def call_chatbot_api(user_message: str):
    url = "https://c44c2f38895b.ngrok-free.app/chat"  

    payload = {"message": user_message, "session_id": "session12345"}
    headers = {"Content-Type": "application/json"}

    response = requests.post(url, json=payload, headers=headers)

    if response.status_code == 200:
        return response.json().get("answer", "Pas de réponse reçue.")
    else:
        return f"Erreur API ({response.status_code}) : {response.text}"

if __name__ == "__main__":
    # Exemple d'utilisation
    cv_path = "C:\\Users\\sebas\\OneDrive\\Bureau\\M5-HETIC\\Agents\\Projet Scrappy Offres\\job-scrape-engine\\CV Sebastian Data.pdf"  # Remplacez par le chemin de votre CV
    analysis = read_cv(cv_path)
    
    #reponse = call_chatbot_api("Salut")
    #print(reponse)
    #print(analysis)